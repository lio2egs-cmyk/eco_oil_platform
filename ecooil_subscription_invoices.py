# -*- coding: utf-8 -*-
"""
Eco-Oil — Limor's subscription invoices for finance (Ravit), automated (14/10/2026).

Limor pays a few software subscriptions (Claude, ChatGPT, Railway) with her own
card for the company and hands the invoices to Ravit for reimbursement. She used
to download the PDFs by hand, file them per supplier under
    Z:\\Eco_General\\משרד_כולל\\לימור_משרד\\קורס AI\\<supplier>
and mail them — and months were missed.

What this script does, daily from Task Scheduler on Limor's PC:

    1. FETCH  — reads the two mailboxes read-only over IMAP (app passwords in
       .env): Anthropic receipts arrive in the private Gmail, Railway receipts
       in office@. Every receipt mail carries Invoice-*.pdf + Receipt-*.pdf;
       a pair not yet filed is saved into the supplier folder on Z:, named the
       way Limor names them ("Invoice-GT2HTDPI-0009_ספטמבר 26.pdf").
       ChatGPT sends no mail: Limor downloads it herself into "צ'אט GPT".
    2. SCAN   — reads every PDF in the three supplier folders (pypdf) and
       extracts invoice number, date paid, amount, description. The folders
       are the only source of truth; the ledger (C:) only remembers what was
       already submitted to Ravit and caches parsed values.
    3. SUBMIT — from the 1st of a month on, the previous month's package is
       sent once all three suppliers have a paid invoice in it (nothing partial
       — Limor's rule): a bordered RTL Excel summary (USD, Bank of Israel
       representative rate on the payment day, ILS) + the PDFs, to Ravit
       (ksafim@) cc Yoav + office. Items of earlier months that were never
       submitted ride along as "השלמות". If something is missing (usually the
       ChatGPT download), a short reminder goes to office@ instead, and the
       package goes out on the first daily run after the gap is filled.

History rule (Limor, 11/10): a file that was already in the folder for a month
up to SUBMITTED_THROUGH counts as submitted and paid. Later ones are pending.

Usage:
    ecooil_subscription_invoices.py              normal daily run
    ecooil_subscription_invoices.py --dry-run    nothing written to Z:, no mail,
                                                 no ledger change — the package
                                                 and mail text go to preview\\
    ecooil_subscription_invoices.py --date 2026-11-01   pretend today is that date
    ecooil_subscription_invoices.py --no-mail    fetch + file, build the package,
                                                 skip sending (stays pending)
"""
import argparse
import email
import email.utils
import imaplib
import json
import os
import re
import shutil
import smtplib
import sys
import traceback
import urllib.request
from datetime import date, datetime, timedelta
from email import encoders
from email.header import Header, decode_header, make_header
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.utils import formataddr

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pypdf import PdfReader

# ---------------------------------------------------------------- settings ---
ENV_FILE = r"C:\eco_oil_platform_git\.env"
Z_BASE = r"Z:\Eco_General\משרד_כולל\לימור_משרד\קורס AI"
SUBMISSIONS_DIR = os.path.join(Z_BASE, "הגשות_לרוית")
WORK_DIR = r"C:\eco_oil_portal\subscriptions"
LEDGER = os.path.join(WORK_DIR, "ledger.json")
LOG_FILE = os.path.join(WORK_DIR, "log.txt")
PREVIEW_DIR = os.path.join(WORK_DIR, "preview")
PACKAGES_DIR = os.path.join(WORK_DIR, "packages")

START_MONTH = (2026, 6)          # Limor: submissions start from June 2026
SUBMITTED_THROUGH = (2026, 8)    # files already on Z: for months up to this = submitted & paid
REMINDER_EVERY_DAYS = 7

MAIL_TO = ["ksafim@eco-oil.co.il"]
MAIL_CC = ["yoav@eco-oil.co.il", "office@eco-oil.co.il"]
OFFICE = "office@eco-oil.co.il"

IMAP_HOST = "imap.gmail.com"
ALL_MAIL = "[Gmail]/All Mail"

SUPPLIERS = {
    "claude": {
        "name": "Claude (Anthropic)",
        "folder": "claude_code",
        "mailbox": "gmail",
        "sender": "invoice+statements@mail.anthropic.com",
        "manual": False,
    },
    "chatgpt": {
        "name": "ChatGPT (OpenAI)",
        "folder": "צ'אט GPT",
        "mailbox": None,
        "sender": None,
        "manual": True,
        "manual_hint": "להוריד מ-chatgpt.com/settings/billing (Transaction history → Open invoice → Download) "
                       "ולשמור בתיקייה 'צ'אט GPT' על הכונן המרכזי",
    },
    "railway": {
        "name": "Railway",
        "folder": "railway",
        "mailbox": "office",
        "sender": "Railway Corporation",
        "manual": False,
    },
}
MAILBOXES = {
    "gmail": ("lio2egs@gmail.com", "SUBS_GMAIL_APP_PASSWORD"),
    "office": ("office@eco-oil.co.il", "SUBS_OFFICE_APP_PASSWORD"),
}

HEB_MONTHS = ["ינואר", "פברואר", "מרץ", "אפריל", "מאי", "יוני",
              "יולי", "אוגוסט", "ספטמבר", "אוקטובר", "נובמבר", "דצמבר"]
IMAP_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
EN_MONTHS = {m: i + 1 for i, m in enumerate(
    ["January", "February", "March", "April", "May", "June", "July", "August",
     "September", "October", "November", "December"])}

INVOICE_RE = re.compile(r"Invoice-([A-Z0-9]{6,10})-(\d{4})", re.I)
RECEIPT_RE = re.compile(r"Receipt-(\d{4}-\d{4}(?:-\d{4})?)", re.I)

DRY = False
TODAY = date.today()


# ------------------------------------------------------------------ helpers --
def log(msg):
    os.makedirs(WORK_DIR, exist_ok=True)
    line = f"{datetime.now():%Y-%m-%d %H:%M:%S} {msg}"
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(line + "\n")
    try:
        print(line)
    except Exception:
        pass


def load_env():
    env = {}
    if os.path.exists(ENV_FILE):
        with open(ENV_FILE, encoding="utf-8", errors="replace") as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    env[k.strip()] = v.strip().strip('"').strip("'")
    return env


def load_ledger():
    if os.path.exists(LEDGER):
        with open(LEDGER, encoding="utf-8") as f:
            return json.load(f)
    return {"items": {}, "packages": {}, "pdf_cache": {}, "rates": {}, "reminders": {}, "pending_mail": []}


def save_ledger(ledger):
    if DRY:
        return
    os.makedirs(WORK_DIR, exist_ok=True)
    tmp = LEDGER + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(ledger, f, ensure_ascii=False, indent=1)
    os.replace(tmp, LEDGER)


def month_key(y, m):
    return f"{y:04d}-{m:02d}"


def month_label(y, m):
    return f"{HEB_MONTHS[m - 1]} {y}"


def file_suffix(d):
    """Limor's naming: '_ספטמבר 26'."""
    return f"_{HEB_MONTHS[d.month - 1]} {d.year % 100:02d}"


def prev_month(d):
    y, m = (d.year, d.month - 1) if d.month > 1 else (d.year - 1, 12)
    return y, m


def ym_le(a, b):
    return (a[0], a[1]) <= (b[0], b[1])


def hdr(msg, k):
    try:
        return str(make_header(decode_header(msg.get(k, ""))))
    except Exception:
        return msg.get(k, "")


def parse_en_date(s):
    m = re.search(r"([A-Z][a-z]+) (\d{1,2}), (\d{4})", s)
    if not m or m.group(1) not in EN_MONTHS:
        return None
    return date(int(m.group(3)), EN_MONTHS[m.group(1)], int(m.group(2)))


# ---------------------------------------------------------------- PDF scan ---
def parse_pdf(path):
    """Invoice number, date paid / issued, amount, description from a Stripe PDF."""
    text = "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
    flat = re.sub(r"[ \t]+", " ", text)
    info = {"kind": None, "invoice": None, "receipt": None, "date": None, "amount": None, "desc": None}
    # Stripe joins the parts with a non-breaking hyphen (U+2011); pypdf may hand it
    # back as that, as a space or as an ordinary hyphen — accept any one character.
    m = re.search(r"Invoice number ([A-Z0-9]{6,10})\W(\d{4})", flat)
    if m:
        info["invoice"] = f"{m.group(1)}-{m.group(2)}"
    m = re.search(r"Receipt number ((?:\d{4}\W){1,2}\d{4})", flat)
    if m:
        info["receipt"] = "-".join(re.findall(r"\d{4}", m.group(1)))
        info["kind"] = "receipt"
    elif info["invoice"]:
        info["kind"] = "invoice"
    m = re.search(r"Date paid (.+)", flat) or re.search(r"Date of issue (.+)", flat)
    if m:
        info["date"] = parse_en_date(m.group(1))
    m = re.search(r"\$([\d,]+\.\d{2}) (?:paid on|USD due)", flat)
    if m:
        info["amount"] = float(m.group(1).replace(",", ""))
    m = re.search(r"Description Qty Unit price Amount\n(.+?)\n(?:Subtotal|Total)", flat, re.S)
    if m:
        lines = [l.strip() for l in m.group(1).splitlines() if l.strip()]
        # the plan line ("Max plan - 20x", "ChatGPT Plus Subscription", "Hobby plan"),
        # not a usage line ("Agent Usage")
        plan = [l for l in lines if re.search(r"plan|subscription|pro\b|max\b", l, re.I)
                and not re.search(r"^\d|\$", l)]
        info["desc"] = (plan or lines)[0] if lines else None
    return info


def dry_folder(key):
    """Where a dry run files what it would have saved on Z:."""
    return os.path.join(PREVIEW_DIR, "fetched", SUPPLIERS[key]["folder"])


def scan_folders(ledger):
    """Every invoice/receipt PDF in the three supplier folders → items by invoice number."""
    cache = ledger.setdefault("pdf_cache", {})
    items = {}
    for key, sup in SUPPLIERS.items():
        folder = os.path.join(Z_BASE, sup["folder"])
        if not os.path.isdir(folder):
            log(f"folder missing: {folder}")
            continue
        listing = [(folder, fn) for fn in sorted(os.listdir(folder))]
        if DRY and os.path.isdir(dry_folder(key)):
            listing += [(dry_folder(key), fn) for fn in sorted(os.listdir(dry_folder(key)))]
        for folder_, fn in listing:
            if not fn.lower().endswith(".pdf"):
                continue
            path = os.path.join(folder_, fn)
            st = os.stat(path)
            ck = f"{key}/{fn}"
            c = cache.get(ck)
            if not c or c.get("size") != st.st_size or c.get("mtime") != int(st.st_mtime):
                try:
                    info = parse_pdf(path)
                except Exception as exc:
                    log(f"pdf unreadable: {ck} — {exc}")
                    info = {"kind": None, "invoice": None, "receipt": None, "date": None, "amount": None, "desc": None}
                c = {"size": st.st_size, "mtime": int(st.st_mtime),
                     "kind": info["kind"], "invoice": info["invoice"], "receipt": info["receipt"],
                     "date": info["date"].isoformat() if info["date"] else None,
                     "amount": info["amount"], "desc": info["desc"]}
                cache[ck] = c
            if not c["invoice"] or not c["date"]:
                continue
            d = date.fromisoformat(c["date"])
            if not ym_le(START_MONTH, (d.year, d.month)):
                continue
            it = items.setdefault(c["invoice"], {
                "supplier": key, "invoice": c["invoice"], "date": c["date"], "amount": c["amount"],
                "desc": c["desc"], "files": {}})
            it["files"][c["kind"] or "other"] = path
            if c["kind"] == "receipt":      # the receipt carries the real paid date
                it["date"], it["amount"] = c["date"], c["amount"] or it["amount"]
            if c["desc"] and not it["desc"]:
                it["desc"] = c["desc"]
    return items


# -------------------------------------------------------------------- IMAP ---
def fetch_receipts(env, ledger):
    """Save unfiled Invoice+Receipt pairs from the receipt mails into the Z: folders."""
    saved = []
    for key, sup in SUPPLIERS.items():
        if sup["manual"]:
            continue
        user, pwkey = MAILBOXES[sup["mailbox"]]
        pw = env.get(pwkey, "")
        if not pw:
            log(f"{key}: no app password ({pwkey}) in .env — fetch skipped")
            continue
        folder = os.path.join(Z_BASE, sup["folder"])
        existing = " ".join(os.listdir(folder)).lower() if os.path.isdir(folder) else ""
        try:
            m = imaplib.IMAP4_SSL(IMAP_HOST, 993, timeout=60)
            m.login(user, pw)
            m.select(f'"{ALL_MAIL}"', readonly=True)
            since = date(START_MONTH[0], START_MONTH[1], 1) - timedelta(days=3)
            since_s = f"{since.day:02d}-{IMAP_MONTHS[since.month - 1]}-{since.year}"
            typ, data = m.search(None, f'(FROM "{sup["sender"]}" SUBJECT "receipt" SINCE {since_s})')
            ids = data[0].split() if data and data[0] else []
            log(f"{key}: {len(ids)} receipt mails in {user}")
            for i in ids:
                typ, md = m.fetch(i, "(RFC822)")
                msg = email.message_from_bytes(md[0][1])
                pdfs = {}
                for p in msg.walk():
                    fn = p.get_filename()
                    if fn and fn.lower().endswith(".pdf"):
                        try:
                            fn = str(make_header(decode_header(fn)))
                        except Exception:
                            pass
                        pdfs[fn] = p.get_payload(decode=True)
                inv = next((INVOICE_RE.search(fn) for fn in pdfs if INVOICE_RE.search(fn)), None)
                if not inv:
                    log(f"{key}: mail without Invoice pdf skipped: {hdr(msg, 'Subject')}")
                    continue
                inv_no = f"{inv.group(1)}-{inv.group(2)}".lower()
                if inv_no in existing:
                    continue
                mp = re.search(r"Paid ([A-Z][a-z]+ \d{1,2}, \d{4})", _plain_body(msg))
                paid = (parse_en_date(mp.group(1)) if mp else None) or \
                    email.utils.parsedate_to_datetime(msg["Date"]).date()
                for fn, blob in pdfs.items():
                    base = os.path.splitext(fn)[0]
                    target_name = base + file_suffix(paid) + ".pdf"
                    target = os.path.join(dry_folder(key) if DRY else folder, target_name)
                    os.makedirs(os.path.dirname(target), exist_ok=True)
                    if os.path.exists(target):
                        continue
                    with open(target, "wb") as f:
                        f.write(blob)
                    saved.append(target)
                    log(f"{key}: filed {target_name}" + (" (preview only)" if DRY else ""))
                existing += " " + inv_no
            m.logout()
        except Exception as exc:
            log(f"{key}: fetch failed — {exc}")
    return saved


def _plain_body(msg):
    for p in msg.walk():
        if p.get_content_type() == "text/plain" and not p.get_filename():
            try:
                return (p.get_payload(decode=True) or b"").decode(p.get_content_charset() or "utf-8", "replace")
            except Exception:
                return ""
    return ""


# ---------------------------------------------------------- exchange rate ---
def boi_rate(ledger, d):
    """Bank of Israel representative USD/ILS rate on day d, or the last one before it."""
    rates = ledger.setdefault("rates", {})
    checked = ledger.setdefault("rates_checked", {})   # days BOI was asked about and had no rate

    def lookup():
        for back in range(0, 12):
            k = (d - timedelta(days=back)).isoformat()
            if k in rates:
                return rates[k], k
            if k not in checked:
                return None                           # an unasked day: must fetch before falling back
        return None

    found = lookup()
    if found:
        return found
    start = (d - timedelta(days=12)).isoformat()
    url = ("https://edge.boi.gov.il/FusionEdgeServer/sdmx/v2/data/dataflow/BOI.STATISTICS/EXR/1.0/"
           f"RER_USD_ILS?startperiod={start}&endperiod={d.isoformat()}&format=csv")
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=30) as r:
        body = r.read().decode("utf-8", "replace")
    got = {}
    for line in body.splitlines()[1:]:
        cols = line.split(",")
        if len(cols) >= 14 and re.fullmatch(r"\d{4}-\d{2}-\d{2}", cols[12]):
            got[cols[12]] = float(cols[13])
    rates.update(got)
    day = date.fromisoformat(start)
    while day <= d and day < date.today():          # today's rate is published only in the afternoon
        if day.isoformat() not in rates:
            checked[day.isoformat()] = True
        day += timedelta(days=1)
    found = lookup()
    if found:
        return found
    raise RuntimeError(f"no BOI rate near {d}")


# ------------------------------------------------------------------ Excel ---
def build_summary(path, title, rows, total_usd, total_ils):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "חשבוניות מנויים"
    ws.sheet_view.rightToLeft = True
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    total_fill = PatternFill("solid", fgColor="DDEBF7")
    ws.merge_cells("A1:H1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="right")
    ws.merge_cells("A2:H2")
    ws["A2"] = ("מנויים ששולמו בכרטיס האשראי הפרטי של לימור טואג עבור אקו אויל. "
                "השער = השער היציג של בנק ישראל ביום החיוב (או האחרון שפורסם לפניו).")
    ws["A2"].alignment = Alignment(horizontal="right", wrap_text=True)
    ws.row_dimensions[2].height = 32
    headers = ["ספק", "תיאור", "מס' חשבונית", "תאריך חיוב", "סכום $", "שער יציג", "ליום", "סכום ₪"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    r = 5
    for row in rows:
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            cell.alignment = Alignment(horizontal="center" if c != 2 else "right", vertical="center")
        ws.cell(row=r, column=5).number_format = "#,##0.00"
        ws.cell(row=r, column=6).number_format = "0.000"
        ws.cell(row=r, column=8).number_format = "#,##0.00"
        r += 1
    ws.cell(row=r, column=1, value="סה\"כ").font = Font(bold=True)
    ws.cell(row=r, column=5, value=total_usd).number_format = "#,##0.00"
    ws.cell(row=r, column=8, value=total_ils).number_format = "#,##0.00"
    for c in range(1, 9):
        cell = ws.cell(row=r, column=c)
        cell.border = border
        cell.fill = total_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for col, w in zip("ABCDEFGH", [18, 30, 18, 13, 11, 10, 12, 12]):
        ws.column_dimensions[col].width = w
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


# -------------------------------------------------------------------- mail ---
def send_mail(env, subject, body, to, cc=(), attachments=()):
    host, user, pw = env.get("MAIL_HOST"), env.get("MAIL_USERNAME"), env.get("MAIL_PASSWORD")
    from_addr = env.get("MAIL_FROM_ADDRESS") or user
    from_name = env.get("MAIL_FROM_NAME", "")
    if not (host and user and pw and from_addr):
        raise RuntimeError("mail channel not configured in .env (MAIL_*)")
    msg = MIMEMultipart("mixed")
    msg["From"] = formataddr((str(Header(from_name, "utf-8")), from_addr)) if from_name else from_addr
    msg["To"] = ", ".join(to)
    if cc:
        msg["Cc"] = ", ".join(cc)
    msg["Subject"] = Header(subject, "utf-8")
    msg.attach(MIMEText(body, "plain", "utf-8"))
    for path in attachments:
        sub = "pdf" if path.lower().endswith(".pdf") else "vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        with open(path, "rb") as f:
            part = MIMEBase("application", sub)
            part.set_payload(f.read())
        encoders.encode_base64(part)
        part.add_header("Content-Disposition", "attachment", filename=("utf-8", "", os.path.basename(path)))
        msg.attach(part)
    with smtplib.SMTP(host, int(env.get("MAIL_PORT", "587")), timeout=30) as s:
        s.ehlo()
        s.starttls()
        s.ehlo()
        s.login(user, pw)
        s.sendmail(from_addr, list(to) + list(cc), msg.as_string())


def deliver(env, ledger, subject, body, to, cc=(), attachments=(), no_mail=False):
    """Send now; if the channel fails, keep the mail pending for the next run."""
    preview_text = f"To: {to}\nCc: {cc}\nSubject: {subject}\nAttachments: {[os.path.basename(a) for a in attachments]}\n\n{body}"
    if DRY or no_mail:
        os.makedirs(PREVIEW_DIR, exist_ok=True)
        fn = os.path.join(PREVIEW_DIR, f"mail_{TODAY:%Y%m%d}_{re.sub(r'[^\w]+', '_', subject)[:40]}.txt")
        with open(fn, "w", encoding="utf-8") as f:
            f.write(preview_text)
        log(f"mail NOT sent ({'dry-run' if DRY else 'no-mail'}): {subject} → {fn}")
        return False
    try:
        send_mail(env, subject, body, to, cc, attachments)
        log(f"mail sent to {to} cc {cc}: {subject}")
        return True
    except Exception as exc:
        log(f"mail failed ({exc}); kept pending: {subject}")
        ledger.setdefault("pending_mail", []).append(
            {"subject": subject, "body": body, "to": list(to), "cc": list(cc), "attachments": list(attachments)})
        return False


def send_pending(env, ledger):
    still = []
    for p in ledger.get("pending_mail", []):
        try:
            send_mail(env, p["subject"], p["body"], p["to"], p.get("cc", ()), p.get("attachments", ()))
            log(f"pending mail sent: {p['subject']}")
        except Exception as exc:
            log(f"pending mail still failing: {p['subject']} — {exc}")
            still.append(p)
    ledger["pending_mail"] = still


# ---------------------------------------------------------------- package ---
def seed_ledger(ledger, items, saved_now):
    """First sight of an item: a file Limor had already filed for a closed month
    counts as submitted; anything the script itself just fetched is pending."""
    L = ledger.setdefault("items", {})
    saved_now = {os.path.normcase(p) for p in saved_now}
    for inv, it in items.items():
        if inv in L:
            continue
        d = date.fromisoformat(it["date"])
        fetched_now = any(os.path.normcase(p) in saved_now for p in it["files"].values())
        if ym_le((d.year, d.month), SUBMITTED_THROUGH) and not fetched_now:
            L[inv] = {"status": "submitted", "package": "historical", "seen": TODAY.isoformat()}
            log(f"{inv} ({it['supplier']} {it['date']}): already on Z: for a closed month — marked submitted")
        else:
            L[inv] = {"status": "pending", "seen": TODAY.isoformat()}
            log(f"{inv} ({it['supplier']} {it['date']} ${it['amount']}): pending submission")


def month_of(it):
    d = date.fromisoformat(it["date"])
    return (d.year, d.month)


def try_submit(env, ledger, items, no_mail):
    target = prev_month(TODAY)
    if not ym_le(START_MONTH, target):
        return
    tkey = month_key(*target)
    if tkey in ledger.setdefault("packages", {}):
        return                                    # this month's package already went out
    L = ledger["items"]
    pending = [it for inv, it in items.items()
               if L.get(inv, {}).get("status") == "pending" and ym_le(month_of(it), target)]
    missing = [sup["name"] for key, sup in SUPPLIERS.items()
               if not any(it["supplier"] == key and month_of(it) == target for it in items.values())]
    if missing:
        remind(env, ledger, target, missing, pending, no_mail)
        return
    if not pending:
        return
    pending.sort(key=lambda it: (it["date"], it["supplier"]))
    rows, attachments, total_usd, total_ils = [], [], 0.0, 0.0
    lines = []
    for it in pending:
        d = date.fromisoformat(it["date"])
        rate, rate_day = boi_rate(ledger, d)
        ils = round(it["amount"] * rate, 2)
        total_usd += it["amount"]
        total_ils += ils
        tag = "" if month_of(it) == target else f" (השלמה, {month_label(*month_of(it))})"
        rows.append([SUPPLIERS[it["supplier"]]["name"] + tag, it["desc"] or "", it["invoice"],
                     d.strftime("%d/%m/%Y"), it["amount"], rate,
                     date.fromisoformat(rate_day).strftime("%d/%m/%Y"), ils])
        lines.append(f"• {SUPPLIERS[it['supplier']]['name']}{tag}: חשבונית {it['invoice']}, "
                     f"{d:%d/%m/%Y}, {it['amount']:,.2f} $ × {rate:.3f} = {ils:,.2f} ₪")
        for kind in ("invoice", "receipt"):
            if kind in it["files"]:
                attachments.append(it["files"][kind])
    label = month_label(*target)
    title = f"חשבוניות מנויים לחודש {label}"
    if any(month_of(it) != target for it in pending):
        title += " + השלמות"
    out_dir = PREVIEW_DIR if DRY else os.path.join(PACKAGES_DIR, tkey)
    xlsx = os.path.join(out_dir, f"חשבוניות מנויים {label}.xlsx")
    build_summary(xlsx, title, rows, round(total_usd, 2), round(total_ils, 2))
    body = (f"שלום רוית,\n\n"
            f"מצורפות חשבוניות המנויים ששולמו בכרטיס האשראי הפרטי של לימור עבור אקו אויל, "
            f"להחזר — {title}:\n\n" + "\n".join(lines) +
            f"\n\nסה\"כ: {total_usd:,.2f} $ = {total_ils:,.2f} ₪ "
            f"(לפי השער היציג של בנק ישראל ביום החיוב).\n"
            f"הפירוט בקובץ האקסל המצורף; החשבוניות והקבלות (PDF) מצורפות.\n\n"
            f"תודה,\nלימור\n\n(נשלח אוטומטית מהפורטל; לשאלות — office@eco-oil.co.il)")
    subject = f"{title} — לימור טואג, להחזר"
    sent = deliver(env, ledger, subject, body, MAIL_TO, MAIL_CC, [xlsx] + attachments, no_mail)
    if sent:
        if not DRY:
            os.makedirs(SUBMISSIONS_DIR, exist_ok=True)
            shutil.copy2(xlsx, os.path.join(SUBMISSIONS_DIR, os.path.basename(xlsx)))
        for it in pending:
            L[it["invoice"]].update({"status": "submitted", "package": tkey, "sent": TODAY.isoformat()})
        ledger["packages"][tkey] = {"sent": TODAY.isoformat(), "items": [it["invoice"] for it in pending],
                                    "usd": round(total_usd, 2), "ils": round(total_ils, 2), "xlsx": xlsx}
        log(f"package {tkey} sent: {len(pending)} items, ${total_usd:.2f} = ₪{total_ils:.2f}")


def remind(env, ledger, target, missing, pending, no_mail):
    tkey = month_key(*target)
    R = ledger.setdefault("reminders", {})
    last = R.get(tkey)
    if last and (TODAY - date.fromisoformat(last)).days < REMINDER_EVERY_DAYS:
        return
    label = month_label(*target)
    hints = [f"• {name}" + (f": {SUPPLIERS[k]['manual_hint']}" if SUPPLIERS[k]["manual"] else
                            " — הקבלה לא נמצאה במייל; לבדוק אם החיוב בוצע")
             for name in missing for k in SUPPLIERS if SUPPLIERS[k]["name"] == name]
    body = (f"לימור שלום,\n\nהגשת חשבוניות המנויים לחודש {label} לרוית ממתינה: חסר —\n" +
            "\n".join(hints) +
            f"\n\nברגע שהקובץ יהיה בתיקייה, ההגשה תישלח אוטומטית בריצה היומית הבאה (08:30).\n"
            f"מה שכבר מוכן: {len(pending)} פריטים.\n\n(תזכורת אוטומטית; חוזרת כל {REMINDER_EVERY_DAYS} ימים)")
    if deliver(env, ledger, f"חשבוניות מנויים {label}: חסר {', '.join(missing)}", body, [OFFICE], (), (), no_mail):
        R[tkey] = TODAY.isoformat()


# -------------------------------------------------------------------- main ---
def main():
    global DRY, TODAY
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--no-mail", action="store_true")
    ap.add_argument("--date")
    a = ap.parse_args()
    DRY = a.dry_run
    if a.date:
        TODAY = date.fromisoformat(a.date)
    log(f"--- run {'DRY ' if DRY else ''}today={TODAY}")
    env = load_env()
    ledger = load_ledger()
    if not os.path.isdir(Z_BASE):
        log(f"Z: folder not reachable: {Z_BASE}")
        deliver(env, ledger, "חשבוניות מנויים: הכונן המרכזי לא זמין",
                f"הסקריפט לא הגיע ל-{Z_BASE}. ינסה שוב מחר.", [OFFICE], (), (), a.no_mail)
        save_ledger(ledger)
        return 1
    try:
        if not DRY:
            send_pending(env, ledger)
        saved = fetch_receipts(env, ledger)
        items = scan_folders(ledger)
        seed_ledger(ledger, items, saved)
        try_submit(env, ledger, items, a.no_mail)
        save_ledger(ledger)
    except Exception:
        err = traceback.format_exc()
        log("FAILED\n" + err)
        deliver(env, ledger, "חשבוניות מנויים: שגיאה בריצה", err[-3000:], [OFFICE], (), (), a.no_mail)
        save_ledger(ledger)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
