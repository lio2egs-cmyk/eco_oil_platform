# -*- coding: utf-8 -*-
"""פורטל הדיפו — התנועות לטנקו בלי להיכנס לפורטל (בקשת ניר, הכרעת לימור 22/10/2026).

מקור אחד בענן לשני ערוצים:
  1. מייל כל שעתיים — רק כשהיו תנועות בחלון. מתג לחברה ("🕑 עדכון תנועות
     כל שעתיים"); הנמענים = נמעני הדוח היומי, אותה רשימה בדיוק.
  2. API משיכה — GET /depot/api/v1/movements עם מפתח ללקוח, קריאה בלבד.

הצינור: הגשר אצל יעל (field_bridge.py → field_movements.py) מחשב את
התנועות מהקובץ החי באותו חישוב בדיוק כמו הדוח היומי ודוחף את 7 ימי העבודה
האחרונים מיד אחרי שהוא כותב לקובץ (ולפחות כל 5 דקות). השדות = מבנה הדוח
היומי בלבד: מס' נכס, מרכז רווח, סוג פעולה, תאריך ושעה (לימור 22/10).

כלל 20/10 (אין תקלה שקטה): יומן לכל שליחה (DepotMovementMail), יומן לכל
פנייה ל-API (DepotApiCall), ונקודת status שהבקרה במחשב של לימור קוראת —
מחשב אחר מזה שדוחף, כך שגשר שנעצר אצל יעל מתגלה.
"""
import hashlib
import ipaddress
import os
import secrets as _secrets
from datetime import date, datetime, timedelta, timezone
from functools import wraps
from zoneinfo import ZoneInfo

from flask import Blueprint, current_app, jsonify, request

from .auth import depot_admin_required
from .db import (db, AdminActionLog, Client, DepotApiCall, DepotApiKey,
                 DepotBridgeStatus, DepotMovement, DepotMovementMail, User)
from .depot_certs import _folder_client_map, _norm
from .ecooil_bridge import _bearer_token
from .mailer import send_office_email

depot_movements = Blueprint("depot_movements", __name__)

IL_TZ = ZoneInfo("Asia/Jerusalem")
WINDOW_WORKDAYS = 7          # היסטוריה: 7 ימי עבודה אחורה (לימור 22/10) — שבת לא נספרת
SLOT_HOURS = 2               # מייל לכל היותר פעם בחלון של שעתיים (לימור 22/10)
# הגנה מפני קריאה פגומה של הקובץ: דחיפה שהייתה מבטלת יותר מזה בבת אחת נדחית
# כולה (קובץ שנקרא ריק / גיליון חסר) — עדיף עיכוב מאשר "בוטל" מזויף ללקוח.
MAX_CANCEL_ABS = 15
MAX_CANCEL_SHARE = 0.3
# API: הגבלת קצב לכל מפתח
RATE_PER_HOUR = 120
# server_time שמוחזר מוקדם בכמה דקות מרגע הפנייה, כדי שדחיפה שנכתבה בדיוק
# באותו רגע לא תיפול בין שתי משיכות. מחיר: תנועה עלולה לחזור פעמיים —
# הצד של הלקוח מזהה לפי id (כתוב במסמך ההסבר לניר).
SINCE_OVERLAP = timedelta(minutes=5)


# ------------------------------------------------------------------ helpers
def _il_now():
    return datetime.now(timezone.utc).astimezone(IL_TZ)


def window_start(today=None):
    """היום הראשון בחלון: 7 ימי עבודה כולל היום, שבת לא נספרת."""
    d = today or _il_now().date()
    n = 0
    while True:
        if d.weekday() != 5:      # 5 = שבת
            n += 1
            if n == WINDOW_WORKDAYS:
                return d
        d -= timedelta(days=1)


def _slot(il_dt):
    return "%s/%02d" % (il_dt.date().isoformat(), il_dt.hour - il_dt.hour % SLOT_HOURS)


def _il_fmt(utc_dt):
    if not utc_dt:
        return None
    return utc_dt.replace(tzinfo=timezone.utc).astimezone(IL_TZ).strftime("%d/%m/%Y %H:%M")


def _any_bridge_token(f):
    """נקודות הבקרה: גם הגשר של יעל (FIELD) וגם הבקרה אצל לימור (ECOOIL)."""
    @wraps(f)
    def wrapper(*args, **kwargs):
        tok = _bearer_token()
        ok = False
        for env in ("FIELD_BRIDGE_TOKEN", "ECOOIL_BRIDGE_TOKEN"):
            exp = os.environ.get(env)
            if exp and tok and _secrets.compare_digest(tok, exp):
                ok = True
        if not ok:
            return jsonify({"error": "unauthorized"}), 401
        return f(*args, **kwargs)
    return wrapper


def _field_bridge_token(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        exp = os.environ.get("FIELD_BRIDGE_TOKEN")
        tok = _bearer_token()
        if not exp or not tok or not _secrets.compare_digest(tok, exp):
            return jsonify({"error": "unauthorized"}), 401
        return f(*args, **kwargs)
    return wrapper


def _clients_by_folder():
    return _folder_client_map()


def _client_folders(client, fmap=None):
    fmap = fmap if fmap is not None else _clients_by_folder()
    return {k for k, c in fmap.items() if c.id == client.id}


def _rows_for_client(client, q, fmap=None):
    keys = _client_folders(client, fmap)
    return [m for m in q if _norm(m.folder) in keys]


# ------------------------------------------------------------ bridge: push
@depot_movements.route("/depot/portal/bridge/movements", methods=["POST"])
@_field_bridge_token
def bridge_push_movements():
    """הגשר של יעל דוחף את כל התנועות של החלון. upsert לפי עוגן השורה;
    תנועה בחלון שלא הגיעה = בוטלה (שורת רפאים שנמחקה / תאריך שתוקן)."""
    body = request.get_json(silent=True) or {}
    items = body.get("items")
    try:
        win_from = date.fromisoformat(str(body.get("window_from") or "")[:10])
    except ValueError:
        return jsonify(error="window_from required (YYYY-MM-DD)"), 400
    if not isinstance(items, list):
        return jsonify(error="items list required"), 400

    now = datetime.utcnow()
    clean = {}
    for it in items:
        key = str(it.get("key") or "").strip()[:200]
        folder = str(it.get("folder") or "").strip()[:200]
        tank = str(it.get("tank") or "").strip()[:40]
        action = str(it.get("action") or "").strip()[:120]
        try:
            d = date.fromisoformat(str(it.get("date") or "")[:10])
        except ValueError:
            continue
        if not key or not folder or not tank or not action or d < win_from:
            continue
        tm = str(it.get("time") or "").strip()[:5] or None
        clean[key] = dict(folder=folder, tank=tank, action=action, move_date=d,
                          move_time=tm,
                          profit_center=(str(it.get("profit_center") or "").strip()[:100] or None))

    existing = {m.key: m for m in DepotMovement.query.filter(
        DepotMovement.move_date >= win_from).all()}
    # עוגן שכבר קיים מחוץ לחלון (תאריך שתוקן קדימה) — נמצא לפי המפתח
    missing_keys = [k for k in clean if k not in existing]
    if missing_keys:
        for m in DepotMovement.query.filter(DepotMovement.key.in_(missing_keys)).all():
            existing[m.key] = m

    to_cancel = [m for k, m in existing.items()
                 if k not in clean and m.cancelled_at is None and m.move_date >= win_from]
    active = sum(1 for m in existing.values() if m.cancelled_at is None)
    if len(to_cancel) > MAX_CANCEL_ABS and len(to_cancel) > active * MAX_CANCEL_SHARE:
        current_app.logger.error("movements push refused: would cancel %d of %d",
                                 len(to_cancel), active)
        return jsonify(error="refused: would cancel %d of %d active movements — "
                             "file read looks incomplete" % (len(to_cancel), active),
                       refused=True), 409

    # הדחיפה הראשונה אי-פעם (טבלה ריקה): מה שתאריכו לפני היום הוא היסטוריה
    # שכבר דווחה בדוחות היומיים — נרשם "כאילו נראה אתמול", לא ממתין למייל ולא
    # מוגש ב-API כחדשה (לימור 22/10: בלי היסטוריה, מתחילים מ"היום").
    bootstrap = DepotMovement.query.first() is None
    il_today = _il_now().date()
    hist_stamp = (datetime.combine(il_today, datetime.min.time()).replace(tzinfo=IL_TZ)
                  .astimezone(timezone.utc).replace(tzinfo=None)) - timedelta(seconds=1)
    added = changed = cancelled = revived = 0
    for k, v in clean.items():
        m = existing.get(k)
        if m is None:
            hist = bootstrap and v["move_date"] < il_today
            db.session.add(DepotMovement(key=k, first_seen_at=hist_stamp if hist else now,
                                         updated_at=hist_stamp if hist else now,
                                         mail_pending=not hist, ever_mailed=hist, **v))
            added += 1
            continue
        diff = any(getattr(m, f) != v[f] for f in
                   ("folder", "tank", "action", "move_date", "move_time", "profit_center"))
        if m.cancelled_at is not None:
            m.cancelled_at = None
            diff = True
            revived += 1
        if diff:
            for f, val in v.items():
                setattr(m, f, val)
            m.updated_at = now
            m.mail_pending = True
            changed += 1
    for m in to_cancel:
        m.cancelled_at = now
        m.updated_at = now
        # ביטול של תנועה שעוד לא דווחה במייל — אין מה לתקן אצל הלקוח
        m.mail_pending = bool(m.ever_mailed)
        cancelled += 1

    st = db.session.get(DepotBridgeStatus, "movements_push")
    if st is None:
        st = DepotBridgeStatus(key="movements_push")
        db.session.add(st)
    st.updated_at = now
    st.value = "%d items, +%d ~%d -%d" % (len(clean), added, changed, cancelled)
    db.session.commit()

    mailed = []
    try:
        mailed = send_movement_mails()
    except Exception as exc:
        db.session.rollback()
        current_app.logger.error("movement mail failed: %s", exc)
    return jsonify(ok=True, items=len(clean), added=added, changed=changed,
                   cancelled=cancelled, revived=revived, mailed=mailed)


# --------------------------------------------------------------- the mail
def _movement_recipients(client):
    """מתג החברה דלוק → בדיוק נמעני הדוח היומי (לימור 22/10/2026)."""
    if not client.movement_mail:
        return []
    return [u for u in User.query.filter_by(client_id=client.id,
                                            role="eco_depot_client").all()
            if u.email and u.is_active and u.daily_report_mail]


def _mail_content(client_name, new_rows, fix_rows, il_now):
    """נוסח טיוטה — ממתין לאישור יואב לפני הפעלה (פונה ללקוח)."""
    stamp = il_now.strftime("%d/%m/%Y %H:%M")
    subject = f"עדכון תנועות אקו-דיפו — {client_name} — {stamp}"

    def td(v, head=False):
        tag = "th" if head else "td"
        bg = "background:#D9D9D9;" if head else ""
        return (f'<{tag} style="border:1px solid #999;padding:4px 10px;{bg}'
                f'text-align:center">{v if v not in (None, "") else ""}</{tag}>')

    def table(rows, with_state=False):
        heads = ["תאריך", "מרכז רווח", "מס' נכס", "סוג פעולה", "שעה"]
        if with_state:
            heads.append("עדכון")
        out = ['<table style="border-collapse:collapse;direction:rtl">',
               "<tr>" + "".join(td(h, True) for h in heads) + "</tr>"]
        for m, state in rows:
            cells = [m.move_date.strftime("%d/%m/%Y"), m.profit_center, m.tank,
                     m.action, m.move_time]
            if with_state:
                cells.append(state)
            out.append("<tr>" + "".join(td(c) for c in cells) + "</tr>")
        out.append("</table>")
        return "".join(out)

    parts = [f'<div dir="rtl" style="font-family:Arial,sans-serif;font-size:15px;line-height:1.7">',
             "<p>שלום,</p>",
             "<p>להלן התנועות שנרשמו אצלנו מאז העדכון הקודם:</p>"]
    if new_rows:
        parts.append(table(new_rows))
    if fix_rows:
        parts.append("<p style=\"margin-top:16px\">עדכונים לתנועות שדווחו קודם:</p>")
        parts.append(table(fix_rows, with_state=True))
    parts.append("<p style=\"margin-top:16px\">מצורף קובץ אקסל במבנה הדוח היומי, עם כל התנועות "
                 "של היום עד לרגע זה. הדוח היומי המלא ממשיך להגיע כרגיל בכל בוקר.</p>"
                 "<p>בברכה,<br>צוות אקו-דיפו</p></div>")
    lines = ["שלום,", "להלן התנועות שנרשמו אצלנו מאז העדכון הקודם:"]
    for m, state in new_rows + fix_rows:
        lines.append(" | ".join(x for x in (m.move_date.strftime("%d/%m/%Y"), m.profit_center or "",
                                            m.tank, m.action, m.move_time or "", state or "") if x))
    lines += ["", "מצורף קובץ אקסל במבנה הדוח היומי, עם כל התנועות של היום עד לרגע זה.",
              "הדוח היומי המלא ממשיך להגיע כרגיל בכל בוקר.", "בברכה, צוות אקו-דיפו"]
    return subject, "".join(parts), "\n".join(lines)


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NO_MOVES_TEXT = "לא נרשמו תנועות ביום זה"


def day_xlsx(cust_name, day, rows, il_now):
    """קובץ מצטבר ליום אחד — במבנה הדוח היומי (gen_daily_pdf.py): A1 כותרת ("עדכון תנועות ביניים"),
    A2 תאריך + סה"כ, A3 הערה, שורה 5 כותרות (מרכז רווח | מס' נכס | סוג פעולה |
    שעה), מרכז רווח כמספר ושעה כזמן h:mm — כמו שאקסל שומר בקובץ היומי, כדי
    שהפריוריטי של הלקוח יקרא אותו בלי שינוי (לימור + ניר 22/10/2026)."""
    from io import BytesIO
    from datetime import time as _time
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "דוח יומי"
    ws.sheet_view.rightToLeft = True
    # הכותרת בלבד שונה מהדוח היומי, כדי שהעובדים יזהו עדכון ביניים (לימור 22/10);
    # הטבלה (שורה 5 והלאה) זהה — אותה קורא הפריוריטי
    ws["A1"] = f"עדכון תנועות ביניים — {cust_name}"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = f"תאריך: {day.strftime('%d/%m/%Y')}    |    סה\"כ פעולות: {len(rows)}"
    ws["A2"].font = Font(size=11)
    ws["A3"] = (f"הופק אוטומטית בתאריך {il_now.strftime('%d/%m/%Y')} בשעה "
                f"{il_now.strftime('%H:%M')} — עדכון ביניים, פירוט תנועות בלבד")
    ws["A3"].font = Font(size=9, italic=True)
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    heads = ["מרכז רווח", "מס' נכס", "סוג פעולה", "שעה"]
    for i, h in enumerate(heads, start=1):
        c = ws.cell(row=5, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9D9D9")
        c.alignment = Alignment(horizontal="center")
        c.border = border
    r = 5
    for m in rows:
        r += 1
        pc = (m.profit_center or "").strip()
        tm = None
        if m.move_time and len(m.move_time) == 5 and m.move_time[2] == ":":
            try:
                tm = _time(int(m.move_time[:2]), int(m.move_time[3:]))
            except ValueError:
                tm = None
        vals = [int(pc) if pc.isdigit() else (pc or None), m.tank, m.action, tm]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = border
        if tm is not None:
            ws.cell(row=r, column=4).number_format = "h:mm"
    if not rows:
        r = 6
        ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=4)
        ws.cell(row=6, column=1, value=NO_MOVES_TEXT).font = Font(italic=True)
        ws.cell(row=6, column=1).alignment = Alignment(horizontal="center")
        for i in range(1, 5):
            ws.cell(row=6, column=i).border = border
    for col, w in zip("ABCD", (37.75, 13.125, 12, 7)):
        ws.column_dimensions[col].width = w
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _day_attachments(client, rows, il_now, fmap):
    """קובץ מצטבר לכל יום שנגעו בו התנועות שבמייל: כל התנועות הפעילות של
    הלקוח באותו יום, עד עכשיו — אותו סדר כמו בדוח היומי (שעה, ואז מכל)."""
    out = []
    keys = _client_folders(client, fmap)
    for day in sorted({m.move_date for m in rows}):
        day_rows = [m for m in DepotMovement.query.filter(
            DepotMovement.move_date == day, DepotMovement.cancelled_at.is_(None)).all()
            if _norm(m.folder) in keys]
        day_rows.sort(key=lambda m: (m.move_time or "99:99", m.tank))
        cust = day_rows[0].folder if day_rows else client.name
        name = "עדכון_תנועות_%s_%s.xlsx" % (day.strftime("%d-%m-%Y"), il_now.strftime("%H-%M"))
        out.append((name, day_xlsx(cust, day, day_rows, il_now), XLSX_MIME))
    return out


def send_movement_mails(dry_run=False, force_slot=False):
    """לכל לקוח עם נמענים מסומנים ותנועות שממתינות: מייל אחד לכל חלון של
    שעתיים (שעון ישראל). לקוח בלי נמענים — לא נשלח כלום והתנועות נשארות
    ממתינות (ראשונה בתור כשיסומן נמען תהיה רק מה שבחלון ההיסטוריה)."""
    il_now = _il_now()
    slot = _slot(il_now)
    now = datetime.utcnow()
    pending = DepotMovement.query.filter_by(mail_pending=True).order_by(
        DepotMovement.move_date, DepotMovement.move_time, DepotMovement.tank).all()
    if not pending:
        return []
    fmap = _clients_by_folder()
    by_client = {}
    for m in pending:
        c = fmap.get(_norm(m.folder))
        if c is not None:
            by_client.setdefault(c.id, (c, []))[1].append(m)
    out = []
    for cid, (c, rows) in by_client.items():
        recips = _movement_recipients(c)
        item = {"client": c.name, "pending": len(rows), "slot": slot,
                "recipients": [u.email for u in recips]}
        if not recips:
            item["result"] = "אין נמענים מסומנים"
            out.append(item)
            continue
        if not force_slot and DepotMovementMail.query.filter_by(
                client_id=cid, slot=slot).filter(
                DepotMovementMail.result != "כשל בשליחה").first():
            item["result"] = "ממתין לחלון הבא"
            out.append(item)
            continue
        # המייל הראשון ללקוח מתחיל מהיום — לא שולחים 7 ימי היסטוריה כ"חדשות"
        # (ההיסטוריה זמינה ב-API ובדוחות היומיים). מכאן והלאה גם רישום מאוחר
        # של יום קודם מדווח כרגיל.
        if not DepotMovementMail.query.filter_by(client_id=cid).filter(
                DepotMovementMail.result != "כשל בשליחה").first():
            today = il_now.date()
            # "היום" = תאריך התנועה היום, או שנרשמה/תוקנה היום (כמו ב-API)
            day_start = (datetime.combine(today, datetime.min.time()).replace(tzinfo=IL_TZ)
                         .astimezone(timezone.utc).replace(tzinfo=None))
            old = [m for m in rows if m.move_date < today and m.first_seen_at < day_start]
            if not dry_run:
                for m in old:
                    m.mail_pending = False
                    if m.cancelled_at is None:
                        m.ever_mailed = True
            rows = [m for m in rows if m not in old]
            item["skipped_history"] = len(old)
            if not rows:
                if not dry_run:
                    db.session.commit()
                item["result"] = "אין תנועות היום"
                out.append(item)
                continue
        new_rows, fix_rows = [], []
        for m in rows:
            if m.cancelled_at is not None:
                fix_rows.append((m, "בוטלה"))
            elif m.ever_mailed:
                fix_rows.append((m, "עודכנה"))
            else:
                new_rows.append((m, None))
        item.update(n_new=len(new_rows), n_fix=len(fix_rows))
        if dry_run:
            item["result"] = "dry_run"
            out.append(item)
            continue
        subject, html, text = _mail_content(c.name, new_rows, fix_rows, il_now)
        files = _day_attachments(c, rows, il_now, fmap)
        sent, failed = [], []
        for u in recips:
            ok = send_office_email(subject=subject, html=html, text=text, to=u.email,
                                   attachments=files)
            (sent if ok else failed).append(u.email)
        result = "נשלח" if sent and not failed else ("נשלח חלקית" if sent else "כשל בשליחה")
        db.session.add(DepotMovementMail(
            client_id=cid, slot=slot, at=now,
            n_new=len(new_rows),
            n_changed=sum(1 for _, s in fix_rows if s == "עודכנה"),
            n_cancelled=sum(1 for _, s in fix_rows if s == "בוטלה"),
            sent_to=", ".join(sent) or None, failed_to=", ".join(failed) or None,
            result=result))
        if sent:
            for m in rows:
                m.mail_pending = False
                if m.cancelled_at is None:
                    m.ever_mailed = True
        db.session.commit()
        item.update(result=result, sent=sent, failed=failed)
        out.append(item)
    return out


# ------------------------------------------------------ status (בקרה)
@depot_movements.route("/depot/portal/bridge/movements/status", methods=["GET"])
@_any_bridge_token
def bridge_movements_status():
    """לבקרה במחשב של לימור: מתי נדחף לאחרונה, מה ממתין למייל ומאז מתי,
    השליחות האחרונות, ומתי כל מפתח API משך לאחרונה."""
    st = db.session.get(DepotBridgeStatus, "movements_push")
    fmap = _clients_by_folder()
    now = datetime.utcnow()
    waiting = {}
    for m in DepotMovement.query.filter_by(mail_pending=True).all():
        c = fmap.get(_norm(m.folder))
        if c is None:
            continue
        w = waiting.setdefault(c.id, {"client": c.name, "count": 0, "oldest_min": 0,
                                      "has_recipients": bool(_movement_recipients(c))})
        w["count"] += 1
        w["oldest_min"] = max(w["oldest_min"], int((now - m.updated_at).total_seconds() // 60))
    last_mails = [{
        "client": (db.session.get(Client, r.client_id).name if r.client_id else "?"),
        "at": _il_fmt(r.at), "slot": r.slot, "result": r.result,
        "new": r.n_new, "changed": r.n_changed, "cancelled": r.n_cancelled,
        "failed_to": r.failed_to,
    } for r in DepotMovementMail.query.order_by(DepotMovementMail.at.desc()).limit(12)]
    keys = [{
        "client": k.client.name if k.client else "?", "prefix": k.prefix,
        "label": k.label,
        "last_used_min": (int((now - k.last_used_at).total_seconds() // 60)
                          if k.last_used_at else None),
        "created": _il_fmt(k.created_at),
    } for k in DepotApiKey.query.filter(DepotApiKey.revoked_at.is_(None)).all()]
    # התראה רק על חסימה של מפתח אמיתי (המחשב של הלקוח נחסם: כתובת לא מורשית,
    # מפתח מבוטל, עומס). פניות בלי מפתח / עם מפתח לא קיים — רובוטים מהאינטרנט
    # או בדיקה שלנו — נרשמות ביומן ונספרות בנפרד, בלי התראה (לימור 25/10/2026,
    # אחרי התראת שווא מבדיקה בלי מפתח).
    since24 = now - timedelta(hours=24)
    refused = DepotApiCall.query.filter(DepotApiCall.at >= since24, DepotApiCall.status != 200,
                                        DepotApiCall.key_id.isnot(None)).count()
    anon = DepotApiCall.query.filter(DepotApiCall.at >= since24, DepotApiCall.status != 200,
                                     DepotApiCall.key_id.is_(None)).count()
    return jsonify(
        last_push_min=(int((now - st.updated_at).total_seconds() // 60)
                       if st and st.updated_at else None),
        last_push_note=st.value if st else None,
        waiting=list(waiting.values()), last_mails=last_mails,
        api_keys=keys, api_refused_24h=refused, api_anon_refused_24h=anon)


# ------------------------------------------------------------------ API
def _client_ip():
    """הכתובת האמיתית: הרשומה הימנית ב-X-Forwarded-For היא זו שהשרת המתווך
    של Railway הוסיף (את השמאליות הלקוח יכול לזייף). ⚠ לאמת בפריסה הראשונה
    מול פנייה מכתובת ידועה."""
    xff = request.headers.get("X-Forwarded-For", "")
    if xff:
        return xff.split(",")[-1].strip()
    return request.remote_addr or "?"


def _ip_allowed(key, ip):
    allowed = [x.strip() for x in (key.allowed_ips or "").split(",") if x.strip()]
    if not allowed:
        return True
    try:
        addr = ipaddress.ip_address(ip)
    except ValueError:
        return False
    for a in allowed:
        try:
            if addr in ipaddress.ip_network(a, strict=False):
                return True
        except ValueError:
            continue
    return False


def _utc_z(dt):
    """זמן UTC בפורמט ISO עם Z — בלי '+' שנשבר בכתובת לא מקודדת."""
    return dt.replace(microsecond=0).isoformat() + "Z"


def _hash_key(raw):
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _api_fail(status, note, key=None, ip=None, since=None):
    db.session.add(DepotApiCall(key_id=key.id if key else None, at=datetime.utcnow(),
                                ip=(ip or "")[:64], since=(since or "")[:40],
                                rows=None, status=status, note=note[:200]))
    db.session.commit()
    msgs = {401: "invalid or missing API key", 403: "forbidden", 429: "rate limit exceeded",
            400: "bad request"}
    return jsonify(error=msgs.get(status, "error"), detail=note), status


@depot_movements.route("/depot/api/v1/movements", methods=["GET"])
def api_movements():
    """API משיכה ללקוח — קריאה בלבד, רק הנכסים שלו, רק שדות הדוח היומי.

    כותרת:  Authorization: Bearer <מפתח>
    פרמטר: since=<ISO-8601>  — התנועות שנוספו/עודכנו/בוטלו אחרי הרגע הזה.
    בלי היסטוריה (לימור 22/10/2026: "מההתחלה נתחיל בעדכון של היום"): מוגש רק
    "היום" — תנועות שתאריכן היום, או שנרשמו/תוקנו/בוטלו היום (רישום מאוחר של
    אתמול שנכנס לקובץ היום הוא חדשה של היום)."""
    ip = _client_ip()
    since_raw = (request.args.get("since") or "").strip()
    raw = _bearer_token() or request.headers.get("X-API-Key", "").strip() or None
    if not raw:
        return _api_fail(401, "no key", ip=ip, since=since_raw)
    key = DepotApiKey.query.filter_by(key_hash=_hash_key(raw)).first()
    if key is None or key.revoked_at is not None:
        return _api_fail(401, "unknown or revoked key" if key is None else "revoked key",
                         key=key, ip=ip, since=since_raw)
    if not _ip_allowed(key, ip):
        return _api_fail(403, "ip not allowed", key=key, ip=ip, since=since_raw)
    now = datetime.utcnow()
    recent = DepotApiCall.query.filter(DepotApiCall.key_id == key.id,
                                       DepotApiCall.at >= now - timedelta(hours=1)).count()
    if recent >= RATE_PER_HOUR:
        return _api_fail(429, "rate limit", key=key, ip=ip, since=since_raw)

    il_today = _il_now().date()
    day_start = (datetime.combine(il_today, datetime.min.time()).replace(tzinfo=IL_TZ)
                 .astimezone(timezone.utc).replace(tzinfo=None))
    since = None
    if since_raw:
        try:
            # "+" בכתובת שלא קודדה מגיע כרווח — מחזירים אותו
            s = datetime.fromisoformat(since_raw.replace(" ", "+").replace("Z", "+00:00"))
        except ValueError:
            return _api_fail(400, "since is not ISO-8601", key=key, ip=ip, since=since_raw)
        if s.tzinfo is not None:
            s = s.astimezone(timezone.utc).replace(tzinfo=None)
        since = s

    q = DepotMovement.query.filter(db.or_(DepotMovement.move_date >= il_today,
                                          DepotMovement.updated_at >= day_start))
    if since is not None:
        q = q.filter(DepotMovement.updated_at > since)
    rows = _rows_for_client(key.client, q.order_by(DepotMovement.updated_at,
                                                   DepotMovement.id).all())
    server_time = now - SINCE_OVERLAP
    out = [{
        "id": m.id,
        "tank": m.tank,
        "profit_center": m.profit_center,
        "action": m.action,
        "date": m.move_date.isoformat(),
        "time": m.move_time,
        "status": "cancelled" if m.cancelled_at else "active",
        "updated_at": _utc_z(m.updated_at),
    } for m in rows]
    key.last_used_at = now
    db.session.add(DepotApiCall(key_id=key.id, at=now, ip=ip[:64], since=since_raw[:40],
                                rows=len(out), status=200, note=None))
    db.session.commit()
    return jsonify(server_time=_utc_z(server_time),
                   today=il_today.isoformat(),
                   count=len(out), movements=out)


# ------------------------------------------------------ admin: keys
def _admin_log(action, details):
    from .depot_admin import _actor
    db.session.add(AdminActionLog(actor=_actor(), division="eco_depot",
                                  action=action, details=details[:400]))


def _key_json(k):
    return {"id": k.id, "prefix": k.prefix, "label": k.label,
            "allowed_ips": k.allowed_ips or "", "created_at": _il_fmt(k.created_at),
            "created_by": k.created_by, "last_used_at": _il_fmt(k.last_used_at),
            "revoked_at": _il_fmt(k.revoked_at)}


@depot_movements.route("/depot/admin/clients/<int:client_id>/api-keys", methods=["GET"])
@depot_admin_required
def admin_list_keys(client_id):
    keys = DepotApiKey.query.filter_by(client_id=client_id).order_by(
        DepotApiKey.created_at.desc()).all()
    calls = (DepotApiCall.query.join(DepotApiKey, DepotApiCall.key_id == DepotApiKey.id)
             .filter(DepotApiKey.client_id == client_id)
             .order_by(DepotApiCall.at.desc()).limit(20).all())
    return jsonify(keys=[_key_json(k) for k in keys], calls=[{
        "at": _il_fmt(c.at), "ip": c.ip, "rows": c.rows, "status": c.status,
        "note": c.note} for c in calls])


@depot_movements.route("/depot/admin/clients/<int:client_id>/api-keys", methods=["POST"])
@depot_admin_required
def admin_create_key(client_id):
    """מנפיק מפתח חדש. המפתח המלא מוחזר פעם אחת בלבד — נשמרת טביעת אצבע."""
    from .depot_admin import _actor
    client = db.session.get(Client, client_id)
    if client is None or client.division != "eco_depot":
        return jsonify(error="חברה לא נמצאה"), 404
    data = request.get_json(silent=True) or {}
    raw = "edk_" + _secrets.token_urlsafe(32)
    k = DepotApiKey(client_id=client.id, key_hash=_hash_key(raw), prefix=raw[:10],
                    label=(str(data.get("label") or "").strip()[:120] or None),
                    allowed_ips=_clean_ips(data.get("allowed_ips")),
                    created_at=datetime.utcnow(), created_by=_actor())
    db.session.add(k)
    _admin_log("הנפקת מפתח API", f"{client.name}: {k.prefix}… ({k.label or '—'})")
    db.session.commit()
    return jsonify(key=raw, **_key_json(k)), 201


def _clean_ips(v):
    parts = []
    for x in str(v or "").replace(";", ",").split(","):
        x = x.strip()
        if not x:
            continue
        try:
            ipaddress.ip_network(x, strict=False)
        except ValueError:
            continue
        parts.append(x)
    return ", ".join(parts)[:400] or None


@depot_movements.route("/depot/admin/api-keys/<int:key_id>", methods=["PATCH"])
@depot_admin_required
def admin_update_key(key_id):
    """ביטול מפתח (מתג הכיבוי) או עדכון הכתובות המורשות."""
    from .depot_admin import _actor
    k = db.session.get(DepotApiKey, key_id)
    if k is None:
        return jsonify(error="מפתח לא נמצא"), 404
    data = request.get_json(silent=True) or {}
    name = k.client.name if k.client else "?"
    if data.get("revoke") and k.revoked_at is None:
        k.revoked_at = datetime.utcnow()
        k.revoked_by = _actor()
        _admin_log("ביטול מפתח API", f"{name}: {k.prefix}…")
    if "allowed_ips" in data:
        k.allowed_ips = _clean_ips(data.get("allowed_ips"))
        _admin_log("כתובות מורשות למפתח API", f"{name}: {k.prefix}… → {k.allowed_ips or 'מכל כתובת'}")
    db.session.commit()
    return jsonify(**_key_json(k))
