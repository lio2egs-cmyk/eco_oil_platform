# -*- coding: utf-8 -*-
"""
Eco-Oil — the quarterly waste-intake report for the Ministry of Environmental
Protection (דוח קליטת פסולות רבעוני), built automatically (Limor, 24/09/2026).

What Limor did by hand each quarter: filter the ריכוז of the three months to the
hazardous streams (everything except צמחי / סניטרי), copy customer + address,
stream, declared tons and package type into the ministry's table, then go
producer by producer and copy ח.פ., toxin-permit number and the waste number
from the master workbook (מסד) — or from an earlier quarterly report.

This script does the same, read-only on every source:

    ריכוז{year}_pivot.xlsx        month sheets of the quarter → one row per unload
    מסד ...xlsx  'ח.פ.-היתר-תוקף'  ח.פ. / היתר רעלים / waste number per stream
    מסד ...xlsx  'הצהרות'          fallback for producers missing above
    earlier quarterly reports       last fallback: what Limor wrote for that producer before

Fixed per-stream columns (Basel Y/A/H codes, UN risk group, treatment, R/D)
are Yoav's rules, taken from the earlier reports (see STREAM_CONST).
Stream sub-types ("מי שטיפה - בוצה") report as the base stream ("מי שטיפה").

The result is ONE sheet in exactly the ministry layout (4 title rows, 15
columns) saved as  דוחות_קליטת פסולות\<year>\רבעון <n>_<year>.xlsx.
Rows whose producer could not be found anywhere get empty ח.פ./היתר/waste
number cells and are listed in the notice so Limor completes them by hand.

Usage:
    ecooil_quarterly_env_report.py                  last completed quarter, file into Z:
    ecooil_quarterly_env_report.py --quarter 2026-Q2
    ecooil_quarterly_env_report.py --dry-run        write to the preview folder, no notice
    ecooil_quarterly_env_report.py --no-notice
"""
import argparse
import glob
import json
import os
import re
import shutil
import smtplib
import traceback
from datetime import date, datetime
from email.header import Header
from email.mime.text import MIMEText
from email.utils import formataddr

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

# ---------------------------------------------------------------- settings ---
ENV_FILE = r"C:\eco_oil_platform_git\.env"
RIKUZ_BASE = r"Z:\Eco_General\ריכוז חודשי"
MASAD = r"Z:\Eco_General\מסד מלא_הצהרות_היתרים_מובילים.xlsx"
REPORTS_BASE = r"Z:\Eco_General\דוחות_קליטת פסולות"
WORK_DIR = r"C:\eco_oil_portal\quarterly_env"
LOG_FILE = os.path.join(WORK_DIR, "log.txt")
BACKUP_DIR = os.path.join(WORK_DIR, "backups")
PREVIEW_DIR = os.path.join(WORK_DIR, "preview")
OFFICE = "office@eco-oil.co.il"

COMPANY = 'אקו אויל חץ וירומטל בע"מ'
PERMIT = 620203

HEB_MONTHS = ["ינואר", "פברואר", "מארס", "אפריל", "מאי", "יוני",
              "יולי", "אוגוסט", "ספטמבר", "אוקטובר", "נובמבר", "דצמבר"]
QUARTER_NAMES = {1: "ראשון", 2: "שני", 3: "שלישי", 4: "רביעי"}
EXCLUDED_STREAMS = {"צמחי", "סניטרי"}

HEADERS = ["יצרן הפסולת", "רשות מקומית", "ח.פ.", "היתר רעלים", "שם זרם הפסולת",
           "מספר הפסולת", "קוד Y - אמנת באזל", "קוד A - אמנת באזל",
           "סיווג בקטלוג הפסולות האירופי", "קוד H - אמנת באזל",
           'קבוצת סיכון על-פי האו"ם', "מתקן/סוג טיפול",
           "קוד טיפול בפסולות (R/D) - אמנת באזל", "כמות (טון)", "סוג האריזות"]
COL_WIDTHS = {"A": 48.2, "B": 24.0, "C": 10.9, "D": 9.8, "E": 9.0, "F": 9.0, "G": 9.0,
              "H": 9.0, "I": 9.0, "J": 9.0, "K": 9.0, "L": 17.0, "M": 13.6, "N": 8.5, "O": 9.0}

# stream → (Y, A, H, UN group, treatment, R/D)  — Yoav's fixed rules as used in
# every quarterly report of 2025-2026
STREAM_CONST = {
    "מינרלי":   ("Y9",  "A4060", "H12", 9, "מערך צנטריפוגלי",    "R12 / R1- D9"),
    "מזוט":     ("Y9",  "A4060", "H12", 9, "מערך צנטריפוגלי",    "R12 / R1- D9"),
    "אמולסיה":  ("Y9",  "A4060", "H12", 9, "ריאקטור פיזיקו-כימי", "R12 / R1- D9"),
    "בסיס":     ("Y35", "A4090", "H8",  8, "מתקן עשת",           "R12/D9"),
    "מי שטיפה": ("Y35", "A4090", "H8",  8, "מתקן עשת",           "R12/D9"),
    "חומצה":    ("Y34", "A4090", "H8",  8, "מתקן עשת",           "R12/D9"),
}
# columns of the מסד sheet 'ח.פ.-היתר-תוקף' holding the waste number per stream
MASAD_WASTE_COL = {"מינרלי": 6, "אמולסיה": 8, "בסיס": 10, "חומצה": 12, "מי שטיפה": 14, "מזוט": 16}
# typo / spelling variants of stream names seen in the ריכוז
STREAM_ALIASES = {"מנרלי": "מינרלי", "מנירלי": "מינרלי", "מינרלי (מזוט)": "מינרלי",
                  "מי-שטיפה": "מי שטיפה", "מי שטיפה": "מי שטיפה"}


# ------------------------------------------------------------------ helpers --
def log(msg):
    os.makedirs(WORK_DIR, exist_ok=True)
    line = f"{datetime.now():%Y-%m-%d %H:%M:%S} {msg}"
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(line + "\n")
    try:
        print(line)
    except Exception:
        pass


def load_env():
    env = {}
    if os.path.exists(ENV_FILE):
        with open(ENV_FILE, encoding="utf-8", errors="replace") as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    env[k.strip()] = v.strip()
    return env


def s(v):
    """Clean string form of a cell (None → '')."""
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def name_key(v):
    """Loose key for matching producer names across files: no spaces, quotes,
    dots, commas, dashes; the '- ממתין לפירוט' style suffixes stay."""
    return re.sub(r"[\s\"'״׳.,\-()]+", "", s(v)).lower()


def base_stream(v):
    """'מי שטיפה - בוצה' → 'מי שטיפה'; typos normalised."""
    raw = s(v)
    base = re.split(r"\s*-\s*", raw)[0].strip() if raw else ""
    return STREAM_ALIASES.get(base, base)


def clean_code(v):
    """Waste numbers as written for the ministry: '160708 *' → '160708*'."""
    t = s(v).replace(" ", "")
    return t


def as_int(v):
    """ח.פ. / permit numbers: keep ints as ints, numeric strings → int."""
    if v is None or s(v) == "":
        return None
    if isinstance(v, (int, float)):
        return int(v)
    t = s(v).replace(",", "").replace("-", "")
    return int(t) if t.isdigit() else s(v)


def quarter_of(d):
    return (d.month - 1) // 3 + 1


def last_completed_quarter(today):
    q = quarter_of(today)
    return (today.year - 1, 4) if q == 1 else (today.year, q - 1)


# ------------------------------------------------------------------ sources --
def read_rikuz_quarter(year, quarter):
    """All hazardous unloads of the quarter, in ריכוז order. Each item:
    {customer, address, stream_raw, stream, tons, package, month, serial}."""
    path = os.path.join(RIKUZ_BASE, str(year), f"ריכוז{year}_pivot.xlsx")
    if not os.path.exists(path):
        raise FileNotFoundError(f"ריכוז workbook not found: {path}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    events, unknown_streams = [], {}
    try:
        for month in range((quarter - 1) * 3 + 1, quarter * 3 + 1):
            name = f"{HEB_MONTHS[month - 1]}_{year % 100:02d}"
            if name not in wb.sheetnames:
                log(f"month sheet {name} missing in {os.path.basename(path)}")
                continue
            for r in wb[name].iter_rows(min_row=2, max_col=17, values_only=True):
                vals = list(r) + [None] * (17 - len(r))
                if not s(vals[1]) or s(vals[1]) == "קוד_רנדומלי":
                    continue                      # incomplete row (no certificate code)
                stream = base_stream(vals[8])
                if not stream or stream in EXCLUDED_STREAMS:
                    continue
                if stream not in STREAM_CONST:
                    unknown_streams[stream] = unknown_streams.get(stream, 0) + 1
                net = vals[11]
                if net is None and vals[9] is not None:
                    net = (vals[9] or 0) - (vals[10] or 0)
                events.append({
                    "customer": s(vals[5]), "address": s(vals[6]),
                    "stream_raw": s(vals[8]), "stream": stream,
                    "tons": round((net or 0) / 1000.0, 3),
                    "package": re.sub(r"\s*-\s*\d+\s*$", "", s(vals[13])),   # 'קוביות-3' → 'קוביות'
                    "month": month, "serial": vals[0],
                    "date": s(vals[2]),
                })
    finally:
        wb.close()
    return events, unknown_streams


def read_masad():
    """Producer → info from the master workbook (two sheets)."""
    wb = openpyxl.load_workbook(MASAD, read_only=True, data_only=True)
    exact, loose, decl = {}, {}, {}
    try:
        for r in wb["ח.פ.-היתר-תוקף"].iter_rows(min_row=2, values_only=True):
            if not s(r[1]):
                continue
            info = {"hp": as_int(r[4]), "permit": as_int(r[2]),
                    "waste": {st: clean_code(r[c]) for st, c in MASAD_WASTE_COL.items() if len(r) > c and s(r[c])},
                    "src": "מסד"}
            exact.setdefault(s(r[1]), info)
            loose.setdefault(name_key(r[1]), info)
        if "הצהרות" in wb.sheetnames:
            for r in wb["הצהרות"].iter_rows(min_row=2, values_only=True):
                if not s(r[1]):
                    continue
                k = name_key(r[1])
                d = decl.setdefault(k, {"hp": None, "permit": None, "waste": {}, "src": "מסד-הצהרות"})
                d["hp"] = d["hp"] or as_int(r[3])
                d["permit"] = d["permit"] or as_int(r[4])
                st = base_stream(r[6])
                if st and s(r[7]) and st not in d["waste"]:
                    d["waste"][st] = clean_code(r[7])
    finally:
        wb.close()
    return exact, loose, decl


def read_earlier_reports(skip_path):
    """What Limor wrote before for each producer: name_key → info (latest wins)."""
    files = sorted(glob.glob(os.path.join(REPORTS_BASE, "*", "רבעון *_*.xlsx")),
                   key=os.path.getmtime)
    hist = {}
    for f in files:
        if os.path.abspath(f) == os.path.abspath(skip_path) or os.path.basename(f).startswith("~$"):
            continue
        try:
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
            ws = wb.worksheets[0]
            for r in ws.iter_rows(min_row=6, max_col=15, values_only=True):
                if not s(r[0]) or s(r[0]) == HEADERS[0]:
                    continue
                k = name_key(r[0])
                info = hist.setdefault(k, {"hp": None, "permit": None, "waste": {}, "src": "דוח קודם"})
                if as_int(r[2]) is not None:
                    info["hp"] = as_int(r[2])
                if as_int(r[3]) is not None:
                    info["permit"] = as_int(r[3])
                st = base_stream(r[4])
                if st and s(r[5]):
                    info["waste"][st] = clean_code(r[5])
            wb.close()
        except Exception as exc:
            log(f"earlier report unreadable {os.path.basename(f)}: {exc}")
    return hist


def lookup(customer, exact, loose, decl, hist):
    """Producer info by exact name → loose name → declarations sheet → earlier
    reports → unique containment in the מסד names. Returns (info, how)."""
    if customer in exact:
        return exact[customer], "מסד"
    k = name_key(customer)
    if k in loose:
        return loose[k], "מסד (שם דומה)"
    merged = None
    for src in (decl.get(k), hist.get(k)):
        if src:
            merged = merged or {"hp": None, "permit": None, "waste": {}, "src": src["src"]}
            merged["hp"] = merged["hp"] or src["hp"]
            merged["permit"] = merged["permit"] or src["permit"]
            for st, code in src["waste"].items():
                merged["waste"].setdefault(st, code)
    if merged:
        return merged, merged["src"]
    # last resort: a מסד name contained in the ריכוז name (site prefixes such as
    # "מיכל 1209 ב-09, גדות אחסון ...") — only when exactly one matches
    cands = [(mk, info) for mk, info in loose.items() if len(mk) >= 8 and mk in k]
    if len(cands) == 1:
        return cands[0][1], "מסד (הכלה)"
    return None, "לא נמצא"


# ------------------------------------------------------------------- build ---
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_rows(events, exact, loose, decl, hist):
    rows, missing, sources = [], [], {}
    for e in events:
        info, how = lookup(e["customer"], exact, loose, decl, hist)
        sources[how] = sources.get(how, 0) + 1
        const = STREAM_CONST.get(e["stream"], ("", "", "", None, "", ""))
        hp = info["hp"] if info else None
        permit = info["permit"] if info else None
        waste = info["waste"].get(e["stream"], "") if info else ""
        if info is None or hp is None or not waste:
            missing.append((e, how, hp, permit, waste))
        rows.append([e["customer"], e["address"], hp, permit, e["stream"], waste,
                     const[0], const[1], waste, const[2], const[3], const[4], const[5],
                     e["tons"], e["package"]])
    return rows, missing, sources


def write_report(path, year, quarter, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "גיליון1"
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.zoomScale = 90
    ws.append(["שם המפעל", COMPANY])
    ws.append(["רבעון", quarter])
    ws.append(["שנה", year])
    ws.append(["מס' היתר רעלים", PERMIT])
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    ws.row_dimensions[2].height = 18.75
    ws.row_dimensions[5].height = 60
    for col, w in COL_WIDTHS.items():
        ws.column_dimensions[col].width = w
    last = 5 + len(rows)
    for r in range(1, last + 1):
        for c in range(1, 16):
            cell = ws.cell(r, c)
            if r <= 4 and c > 2:
                continue
            cell.border = BORDER
            cell.font = Font(name="Calibri", size=11, bold=(r == 5))
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=(r == 5))
            if c == 14 and r >= 5:
                cell.number_format = "0.000"
    wb.save(path)
    chk = openpyxl.load_workbook(path, read_only=True)
    try:
        n = sum(1 for r in chk.worksheets[0].iter_rows(min_row=6, values_only=True) if s(r[0]))
    finally:
        chk.close()
    if n != len(rows):
        raise RuntimeError(f"verification after save failed: {n} rows written, {len(rows)} expected")


def backup(path):
    if not os.path.exists(path):
        return None
    os.makedirs(BACKUP_DIR, exist_ok=True)
    dst = os.path.join(BACKUP_DIR, f"{os.path.splitext(os.path.basename(path))[0]}_{datetime.now():%Y-%m-%d_%H%M%S}.xlsx")
    shutil.copy2(path, dst)
    return dst


def has_data_rows(path):
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        n = sum(1 for r in wb.worksheets[0].iter_rows(min_row=6, max_col=1, values_only=True) if s(r[0]))
        wb.close()
        return n > 0
    except Exception:
        return True   # unreadable → treat as precious


# ------------------------------------------------------------------ notice ---
def send_notice(env, subject, body):
    host, user, pw = env.get("MAIL_HOST"), env.get("MAIL_USERNAME"), env.get("MAIL_PASSWORD")
    from_addr = env.get("MAIL_FROM_ADDRESS") or user
    from_name = env.get("MAIL_FROM_NAME", "")
    if not (host and user and pw and from_addr):
        raise RuntimeError("mail channel not configured (MAIL_* in .env)")
    msg = MIMEText(body, "plain", "utf-8")
    msg["From"] = formataddr((str(Header(from_name, "utf-8")), from_addr)) if from_name else from_addr
    msg["To"] = OFFICE
    msg["Subject"] = Header(subject, "utf-8")
    with smtplib.SMTP(host, int(env.get("MAIL_PORT", "587")), timeout=30) as srv:
        srv.ehlo(); srv.starttls(); srv.ehlo()
        srv.login(user, pw)
        srv.sendmail(from_addr, [OFFICE], msg.as_string())


# -------------------------------------------------------------------- main ---
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--quarter", help="YYYY-Qn (default: last completed quarter)")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--no-notice", action="store_true")
    args = ap.parse_args()
    if args.quarter:
        m = re.fullmatch(r"(\d{4})-Q([1-4])", args.quarter.strip().upper())
        if not m:
            raise SystemExit("--quarter must look like 2026-Q3")
        year, quarter = int(m.group(1)), int(m.group(2))
    else:
        year, quarter = last_completed_quarter(date.today())
    env = load_env()
    fname = f"רבעון {QUARTER_NAMES[quarter]}_{year}.xlsx"
    target = os.path.join(REPORTS_BASE, str(year), fname)
    log(f"=== {'DRY ' if args.dry_run else ''}quarter {year}-Q{quarter} → {fname}")
    try:
        events, unknown = read_rikuz_quarter(year, quarter)
        exact, loose, decl = read_masad()
        hist = read_earlier_reports(target)
        rows, missing, sources = build_rows(events, exact, loose, decl, hist)
        log(f"events {len(events)}; sources {sources}; incomplete rows {len(missing)}; unknown streams {unknown}")

        if args.dry_run:
            os.makedirs(PREVIEW_DIR, exist_ok=True)
            out = os.path.join(PREVIEW_DIR, fname)
        else:
            os.makedirs(os.path.dirname(target), exist_ok=True)
            out = target
            if os.path.exists(target):
                b = backup(target)
                log(f"existing file backed up: {b}")
                if has_data_rows(target):
                    out = os.path.join(REPORTS_BASE, str(year), f"רבעון {QUARTER_NAMES[quarter]}_{year} (אוטומטי).xlsx")
                    log(f"existing file already has rows — writing beside it: {os.path.basename(out)}")
        write_report(out, year, quarter, rows)
        log(f"written: {out} ({len(rows)} rows, {sum(r[13] for r in rows):,.3f} tons)")

        lines = [f"הדוח הרבעוני למשרד להגנת הסביבה, רבעון {QUARTER_NAMES[quarter]} {year}, מוכן.",
                 f"נמצא ב: {out}", f"{len(rows)} שורות, {sum(r[13] for r in rows):,.1f} טון."]
        if missing:
            lines.append(f"\n{len(missing)} שורות להשלמה ידנית (ח.פ. / היתר / מספר פסולת לא נמצאו במסד ולא בדוחות קודמים):")
            seen = set()
            for e, how, hp, permit, waste in missing:
                k = (e["customer"], e["stream"])
                if k in seen:
                    continue
                seen.add(k)
                what = [w for w, v in (("ח.פ.", hp), ("היתר", permit), ("מספר פסולת", waste)) if not v]
                lines.append(f"  - {e['customer']} | {e['stream']} | חסר: {', '.join(what)}")
        if unknown:
            lines.append("\nזרמים שאין להם קודים קבועים (נכנסו בלי קודי באזל): " + ", ".join(f"{k} ({v})" for k, v in unknown.items()))
        body = "\n".join(lines)
        with open(os.path.join(os.path.dirname(out), "_notice.txt") if args.dry_run else os.path.join(WORK_DIR, f"notice_{year}-Q{quarter}.txt"), "w", encoding="utf-8") as f:
            f.write(body)
        if not args.dry_run and not args.no_notice:
            try:
                send_notice(env, f"דוח רבעוני למשרד להגנת הסביבה — רבעון {QUARTER_NAMES[quarter]} {year} מוכן", body)
                log("notice sent to office")
            except Exception as exc:
                log(f"notice mail failed: {exc}")
    except Exception:
        log("ERROR\n" + traceback.format_exc())
        if not args.dry_run:
            try:
                send_notice(env, "דוח רבעוני למשרד להגנת הסביבה — תקלה", traceback.format_exc())
            except Exception:
                pass
        raise
    log("=== done")


if __name__ == "__main__":
    main()
